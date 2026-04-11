"""
BC SCCR — Formatting rules
"""

import re
from docx.oxml import OxmlElement

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

_FORM_HEADER    = re.compile(r"^Form\s+[\d.]+", re.IGNORECASE)
_RULE_NOTICE    = re.compile(r"^\[Rule\s+\d", re.IGNORECASE)
_NUMBERED_ITEM  = re.compile(r"^\d+[\xa0\s]{1,4}\S")   # "1  The facts..."
_BARE_NUMBER    = re.compile(r"^\d+\.?$")               # "1" or "1."


def _set_align(para, value):
    pPr = para._p.get_or_add_pPr()
    jc = pPr.find(f"{{{W}}}jc")
    if jc is None:
        jc = OxmlElement("w:jc")
        pPr.append(jc)
    jc.set(f"{{{W}}}val", value)


def _strip_leading_number(para):
    """Remove the leading 'N  ' prefix from paragraph runs in-place."""
    prefix = re.match(r"^\d+[\xa0\s]+", para.text)
    if not prefix:
        return
    remaining = len(prefix.group(0))
    for run in para.runs:
        if remaining <= 0:
            break
        rt = run.text
        if len(rt) <= remaining:
            run.text = ""
            remaining -= len(rt)
        else:
            run.text = rt[remaining:]
            remaining = 0


def _restart_list_at(doc, para):
    """Make this paragraph begin a new numbered list starting at 1."""
    from docx.oxml.ns import qn

    try:
        numbering_part = doc.part.numbering_part
    except Exception:
        return
    if numbering_part is None:
        return

    numbering_elem = numbering_part._element
    existing_nums  = numbering_elem.findall(qn("w:num"))
    if not existing_nums:
        return

    # Find the abstractNumId the List Number style uses
    style_pPr = doc.styles["List Number"].element.find(qn("w:pPr"))
    abstract_id = None
    if style_pPr is not None:
        style_numPr = style_pPr.find(qn("w:numPr"))
        if style_numPr is not None:
            nid_elem = style_numPr.find(qn("w:numId"))
            if nid_elem is not None:
                style_nid = nid_elem.get(qn("w:val"))
                for num in existing_nums:
                    if num.get(qn("w:numId")) == style_nid:
                        abs_ref = num.find(qn("w:abstractNumId"))
                        if abs_ref is not None:
                            abstract_id = abs_ref.get(qn("w:val"))
                        break

    if abstract_id is None:
        abs_ref = existing_nums[0].find(qn("w:abstractNumId"))
        if abs_ref is None:
            return
        abstract_id = abs_ref.get(qn("w:val"))

    # Create a new w:num with startOverride = 1
    new_id  = max(int(n.get(qn("w:numId"), 0)) for n in existing_nums) + 1
    new_num = OxmlElement("w:num")
    new_num.set(qn("w:numId"), str(new_id))
    abs_num = OxmlElement("w:abstractNumId")
    abs_num.set(qn("w:val"), abstract_id)
    new_num.append(abs_num)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_ov = OxmlElement("w:startOverride")
    start_ov.set(qn("w:val"), "1")
    override.append(start_ov)
    new_num.append(override)
    numbering_elem.append(new_num)

    # Attach the new numId explicitly to this paragraph
    pPr     = para._p.get_or_add_pPr()
    old_np  = pPr.find(qn("w:numPr"))
    if old_np is not None:
        pPr.remove(old_np)
    numPr   = OxmlElement("w:numPr")
    ilvl    = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numPr.append(ilvl)
    numId_e = OxmlElement("w:numId")
    numId_e.set(qn("w:val"), str(new_id))
    numPr.append(numId_e)
    pPr.append(numPr)


def _rewrite_cell(cell, new_text):
    first = None
    for para in cell.paragraphs:
        for run in para.runs:
            if first is None:
                first = run
                first.text = new_text
            else:
                run.text = ""
        return


def run(doc):
    past_appendix = False
    in_list       = False

    for para in doc.paragraphs:
        t = para.text.strip()

        if t == "Appendix":
            past_appendix = True
        if past_appendix:
            continue

        # Form header → centered
        if _FORM_HEADER.match(t):
            _set_align(para, "center")
            in_list = False

        # All-caps title → centered (bold already set in source)
        elif t.isupper() and len(t) > 3 and not t.startswith("["):
            _set_align(para, "center")
            in_list = False

        # Rule notice line → centered
        elif _RULE_NOTICE.match(t):
            _set_align(para, "center")
            in_list = False

        # Numbered content paragraph → List Number, strip leading number
        elif _NUMBERED_ITEM.match(para.text):
            _strip_leading_number(para)
            para.style = doc.styles["List Number"]
            if not in_list:
                _restart_list_at(doc, para)
            in_list = True

        # Bare number placeholder → blank List Number item
        elif _BARE_NUMBER.match(t):
            for run in para.runs:
                run.text = ""
            para.style = doc.styles["List Number"]
            if not in_list:
                _restart_list_at(doc, para)
            in_list = True

        else:
            in_list = False

    # Dot-only table cells → underscores
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                ct = cell.text.strip()
                if ct and re.match(r"^\.+$", ct):
                    _rewrite_cell(cell, "________________________")
