"""
BC SCCR — Placeholder insertion
"""

import re
from docx.oxml import OxmlElement

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

_FILLER = {
    "set", "out", "if", "any", "the", "a", "an", "of", "in", "to", "and",
    "or", "for", "your", "this", "that", "as", "at", "by", "is", "it", "on",
    "up", "be", "do", "use", "using", "include", "each", "all", "with",
    "from", "one", "other", "such", "where", "applicable",
}

_ADDRESS_LABEL = re.compile(
    r"^((?:defendant|plaintiff|party|petitioner|respondent)[^:]*address for service"
    r"|fax number[^:]*|e-?mail[^:]*)",
    re.IGNORECASE,
)
_FILED_BY = re.compile(r"^Filed by:", re.IGNORECASE)


def _infer(text, counter):
    """Derive a short snake_case name from descriptive text."""
    t = re.sub(r"^\[|\]$", "", text.lower()).strip()
    t = re.sub(r"\(if any\)", "", t)
    t = re.sub(r"address for service", "address", t)
    t = re.sub(r"[^a-z0-9]+", " ", t).strip()
    words = [w for w in t.split() if w not in _FILLER and len(w) > 1][:3]
    if words:
        return "_".join(words)
    counter[0] += 1
    return f"field_{counter[0]}"


def _make_para_elem(text, align=None):
    """Create a w:p XML element with optional alignment."""
    p = OxmlElement("w:p")
    if align:
        pPr = OxmlElement("w:pPr")
        jc = OxmlElement("w:jc")
        jc.set(f"{{{W}}}val", align)
        pPr.append(jc)
        p.append(pPr)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    if text != text.strip():
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    p.append(r)
    return p


def _clear_runs(para):
    p_elem = para._p
    for r in list(p_elem.findall(f"{{{W}}}r")):
        p_elem.remove(r)


def _rewrite_para(para, new_text):
    """Write new_text into first run, clear the rest."""
    first = None
    for run in para.runs:
        if first is None:
            first = run
            first.text = new_text
        else:
            run.text = ""
    if first is None and new_text:
        para.add_run(new_text)


def _rewrite_para_label_bold(para, text):
    """Rewrite paragraph so text up to first ':' is bold, rest is plain."""
    colon = text.find(":")
    if colon == -1:
        _rewrite_para(para, text)
        return
    label = text[:colon + 1]
    value = text[colon + 1:]
    _clear_runs(para)
    r1 = para.add_run(label)
    r1.bold = True
    if value:
        r2 = para.add_run(value)
        r2.bold = False


def _rewrite_cell(cell, new_text):
    for para in cell.paragraphs:
        _rewrite_para(para, new_text)
        return


def _expand_style_of_proceeding(doc):
    """Replace [Style of Proceeding] with the 3-line header block."""
    for para in doc.paragraphs:
        if para.text.strip() == "[Style of Proceeding]":
            p_elem = para._element
            parent = p_elem.getparent()
            idx = list(parent).index(p_elem)
            parent.remove(p_elem)
            for i, (text, align) in enumerate([
                ("File Number:{{file_number}}",    "right"),
                ("Registry:{{registry_location}}", "right"),
                ("{{Style_of_Cause}}",             "center"),
            ]):
                parent.insert(idx + i, _make_para_elem(text, align))
            return


def _transform_text(text, counter):
    """Apply placeholder substitutions to a paragraph text string."""

    # Fixed: filing party dot-fill → always client_name
    text = re.sub(r"\.{4,}\[party\(ies\)\]\.{4,}", "{{client_name}}", text)

    # Fixed: address for service with long bracketed instruction after label
    text = re.sub(
        r"(address for service:\s*)\[.*?\]",
        r"\1{{address}}",
        text,
        flags=re.IGNORECASE,
    )

    # Bare fax label ending with colon and nothing after
    text = re.sub(
        r"^(Fax number[^:]*):(\s*)$",
        lambda m: m.group(1) + ":{{fax_number}}",
        text,
        flags=re.IGNORECASE,
    )

    # Bare email label ending with colon and nothing after
    text = re.sub(
        r"^(E-?mail[^:]*):(\s*)$",
        lambda m: m.group(1) + ":{{email}}",
        text,
        flags=re.IGNORECASE,
    )

    # Mid-sentence dot-fills (surrounded by other text) → underscores
    text = re.sub(r"\.{4,}\[[^\]]+\]\.{4,}", "________________", text)

    return text


def _fix_signature_table(doc):
    """Fix date cell and signature checkbox in tables."""
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()

                # Date cell
                if re.match(r"^Date:", t):
                    _rewrite_cell(cell, "Date: _________________")

                # Signature checkbox → {{lawyer_name}}
                elif re.search(r"\[[\s\xa0]*\].*\[[\s\xa0]*\]", t):
                    after = re.split(r"\[[\s\xa0]*\][^\[]*\[[\s\xa0]*\]", t)[-1].strip()
                    _rewrite_cell(cell, "Signature of {{lawyer_name}} " + after)


def run(doc):
    counter = [0]
    past_appendix = False

    _expand_style_of_proceeding(doc)

    for para in doc.paragraphs:
        if para.text.strip() == "Appendix":
            past_appendix = True
        if past_appendix:
            continue

        new_text = _transform_text(para.text, counter)
        if new_text == para.text:
            continue

        # Address / fax / email lines: label bold, value plain
        if _ADDRESS_LABEL.match(new_text.strip()):
            _rewrite_para_label_bold(para, new_text)

        # Filed by line: label bold, value plain
        elif _FILED_BY.match(new_text.strip()):
            _rewrite_para_label_bold(para, new_text)

        else:
            _rewrite_para(para, new_text)

    _fix_signature_table(doc)
